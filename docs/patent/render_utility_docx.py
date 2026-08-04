#!/usr/bin/env python3
"""Create editable DOCX copies with embedded skill-generated figures."""

from __future__ import annotations

import re
import sys
import os
from pathlib import Path

sys.path.insert(0, os.environ.get("PATENT_DOCX_DEPS", "/tmp/patent_docx_deps"))
from docx import Document  # type: ignore
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output" / "utility_rendered"
SERIF = "Noto Serif CJK SC"
SANS = "Noto Sans CJK SC"


def set_font(run, name=SERIF, size=10.5, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def add_inline(p, text, size=10.5, force_bold=False):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = p.add_run(text[pos:match.start()])
            set_font(run, SERIF, size, force_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = p.add_run(token[2:-2])
            set_font(run, SANS, size, True)
        else:
            run = p.add_run(token[1:-1])
            set_font(run, "Noto Sans Mono CJK SC", size)
        pos = match.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        set_font(run, SERIF, size, force_bold)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = SERIF
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("实用新型专利撰写初稿"), SANS, 8)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.55
    p.paragraph_format.space_after = Pt(3)
    add_inline(p, text, 10.2, bold)
    return p


def add_heading(doc, title, level):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level <= 2 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(11 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(title), SANS, 17 if level == 1 else (13 if level == 2 else 11.5), True)


def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), SANS, 9, True)
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cells[i].paragraphs[0]
            add_inline(p, value, 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build(source: Path):
    doc = Document()
    configure(doc)
    lines = source.read_text(encoding="utf-8").splitlines()
    paragraph = []
    i = 0

    def flush():
        nonlocal paragraph
        if paragraph:
            text = " ".join(x.strip() for x in paragraph)
            add_para(doc, text, bool(re.match(r"^\*\*\d+\.", text)))
            paragraph = []

    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            flush(); i += 1; continue
        if stripped == "---":
            flush(); doc.add_paragraph(); i += 1; continue
        if stripped.startswith(">"):
            flush(); add_para(doc, stripped.lstrip("> ")); i += 1; continue
        heading = re.fullmatch(r"(#{1,3})\s+(.+)", stripped)
        if heading:
            flush(); add_heading(doc, heading.group(2), len(heading.group(1)))
            if len(heading.group(1)) == 1 and heading.group(2) in {"权利要求书", "说明书", "附图"}:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1; continue
        image = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", stripped)
        if image:
            flush()
            alt, rel = image.groups()
            path = (source.parent / rel).resolve()
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Inches(6.1))
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(cap.add_run(alt), SANS, 9.5)
            i += 1; continue
        if stripped.startswith("|") and i + 1 < len(lines):
            sep = [x.strip() for x in lines[i + 1].strip().strip("|").split("|")]
            if sep and all(re.fullmatch(r":?-{3,}:?", x) for x in sep):
                flush(); rows = [[x.strip() for x in stripped.strip("|").split("|")]]; i += 2
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")]); i += 1
                add_table(doc, rows); continue
        paragraph.append(lines[i]); i += 1
    flush()
    target = OUT / f"{source.stem}.docx"
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(target)


def main():
    for name in ["实用新型一_羽毛球动作证据反馈装置.md", "实用新型二_骨骼与球路联合时序识别装置.md"]:
        build(ROOT / name)


if __name__ == "__main__":
    main()
