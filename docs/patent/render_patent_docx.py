#!/usr/bin/env python3
"""Build a portable DOCX with embedded patent figures.

This companion renderer uses python-docx when available.  It is kept separate
from the dependency-free HTML renderer because LibreOffice may preserve HTML
images as external links during DOCX export.
"""

from __future__ import annotations

import os
import re
import sys
from pathlib import Path

sys.path.insert(0, "/tmp/patent_docx_deps")
from docx import Document  # type: ignore
from docx.enum.section import WD_SECTION
from docx.image.image import Image as DocxImage
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "羽毛球运动分析发明专利初稿.md"
TARGET = ROOT / "output" / "羽毛球运动分析发明专利初稿.docx"
FONT = "Noto Serif CJK SC"
SANS = "Noto Sans CJK SC"


def set_run_font(run, name: str = FONT, size: float | None = None, bold: bool | None = None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, fill: str = "F2F2F2"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, SANS if bold else FONT, 8.8, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_inline(paragraph, text: str, size: float = 10.5):
    """Add a small subset of Markdown inline markup to a paragraph."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, FONT, size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, SANS, size, True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Noto Sans Mono CJK SC", size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, FONT, size)


def add_para(doc: Document, text: str, *, claim: bool = False, note: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.65
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Cm(0)  # paragraph numbers are explicit
    if note:
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.right_indent = Cm(0.3)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
    add_inline(p, text, 10.3 if not claim else 10.2)
    if note:
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F7F7F7")
        pPr.append(shd)
    return p


def add_heading(doc: Document, title: str, level: int):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level <= 2 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    set_run_font(run, SANS, 17 if level == 1 else (13 if level == 2 else 11.5), True)
    return p


def page_break(doc: Document):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def parse_table(lines: list[str], start: int):
    def cells(line: str):
        return [part.strip() for part in line.strip().strip("|").split("|")]

    headers = cells(lines[start])
    i = start + 2
    body = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        body.append(cells(lines[i]))
        i += 1
    return headers, body, i


def add_table(doc: Document, headers: list[str], body: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    # python-docx otherwise inherits a very large default grid width in some
    # locale profiles (LibreOffice then refuses to open the resulting DOCX).
    # Set an explicit text-width table and equal columns.
    text_width_cm = 16.5
    column_width_cm = text_width_cm / max(1, len(headers))
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(text_width_cm / 2.54 * 1440)))
    for column in table.columns:
        column.width = Cm(column_width_cm)
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].width = Cm(column_width_cm)
        set_cell_text(table.rows[0].cells[idx], header, True)
        shade_cell(table.rows[0].cells[idx])
    for row in body:
        cells = table.add_row().cells
        for idx in range(len(headers)):
            cells[idx].width = Cm(column_width_cm)
            set_cell_text(cells[idx], row[idx] if idx < len(row) else "")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, alt: str, source: str):
    page_break(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    path = (ROOT / source).resolve()
    if not path.is_file():
        # Markdown paths are relative to docs/patent; this branch is useful if
        # a hand-edited draft points to a source-local path.
        path = (ROOT / source.lstrip("./")).resolve()
    run = p.add_run()
    # Keep figures inside the printable A4 area while preserving their aspect
    # ratio.  This matters for tall UI screenshots (for example, the new
    # ShuttleVision workbench figure 8), which would otherwise retain a
    # 16-cm width and extend far beyond the page.
    image = DocxImage.from_file(str(path))
    max_width_cm = 16.0
    max_height_cm = 22.5
    aspect = image.height / image.width if image.width else 1.0
    if aspect > max_height_cm / max_width_cm:
        run.add_picture(str(path), height=Cm(max_height_cm))
    else:
        run.add_picture(str(path), width=Cm(max_width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    r = cap.add_run(alt)
    set_run_font(r, SANS, 9.8)


def configure_document(doc: Document):
    section = doc.sections[0]
    # python-docx's Cm() argument is centimetres (not millimetres).
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    # Keep the footer deliberately simple.  Some older LibreOffice builds
    # reject a DOCX containing a hand-built field element even though Word
    # accepts it; page numbering can be added by the applicant in Word.
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("专利撰写初稿")
    set_run_font(run, SANS, 8)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    max_lines = int(os.environ.get("PATENT_DOCX_MAX_LINES", "0") or 0)
    if max_lines > 0:
        lines = lines[:max_lines]
    doc = Document()
    configure_document(doc)
    paragraph: list[str] = []
    in_formula = False
    formula: list[str] = []

    def flush():
        nonlocal paragraph
        if not paragraph:
            return
        text = " ".join(part.strip() for part in paragraph)
        if text.startswith("> "):
            text = text[2:]
            add_para(doc, text, note=True)
        else:
            add_para(doc, text, claim=bool(re.match(r"\*\*\d+\.", text)))
        paragraph = []

    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped == "$$":
            flush()
            if in_formula:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(5)
                r = p.add_run("\n".join(formula))
                set_run_font(r, FONT, 10.5)
                formula = []
                in_formula = False
            else:
                in_formula = True
            i += 1
            continue
        if in_formula:
            formula.append(line)
            i += 1
            continue

        image_match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", stripped)
        if image_match:
            flush()
            add_figure(doc, *image_match.groups())
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines):
            separator = [x.strip() for x in lines[i + 1].strip().strip("|").split("|")]
            if separator and all(re.fullmatch(r":?-{3,}:?", x) for x in separator):
                flush()
                headers, body, i = parse_table(lines, i)
                add_table(doc, headers, body)
                continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            flush()
            level = len(heading.group(1))
            title = heading.group(2)
            if first_heading and level == 1:
                # Draft title page.
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(55)
                p.paragraph_format.space_after = Pt(20)
                r = p.add_run(title)
                set_run_font(r, SANS, 20, True)
                first_heading = False
            else:
                if title in {"权利要求书", "说明书", "说明书附图", "撰写校核附录（不纳入正式申请文件）"}:
                    page_break(doc)
                add_heading(doc, title, level)
            i += 1
            continue

        if stripped == "---":
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "777777")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        if stripped.startswith(">"):
            flush()
            paragraph.append(stripped)
            flush()
            i += 1
            continue

        list_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if list_match:
            flush()
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.line_spacing = 1.5
            add_inline(p, list_match.group(1), 10)
            i += 1
            continue

        if not stripped:
            flush()
        else:
            paragraph.append(line)
        i += 1
    flush()
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    build()
